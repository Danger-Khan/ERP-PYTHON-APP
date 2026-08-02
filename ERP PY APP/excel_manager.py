import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data')

def clean_val(val):
    if val is None or str(val).strip().lower() in ["none", "nan", "null"]:
        return ""
    return str(val).strip()

class ExcelManager:
    def __init__(self, data_dir=DATA_DIR):
        self.data_dir = data_dir
        os.makedirs(self.data_dir, exist_ok=True)
        self._file_mtimes = {}

    def get_filepath(self, filename):
        """Checks main script directory first, then falls back to data/ folder."""
        base_dir = os.path.dirname(os.path.abspath(__file__))
        root_filepath = os.path.join(base_dir, filename)
        data_filepath = os.path.join(self.data_dir, filename)
        
        if os.path.exists(root_filepath):
            return root_filepath
        return data_filepath

    # ==========================================
    # STYLING ENGINE FOR EXCEL FILES
    # ==========================================
    def style_worksheet(self, ws, headers):
        """Applies headers styling, grid borders, and auto-adjusts column widths."""
        header_fill = PatternFill(start_color="007AFF", end_color="007AFF", fill_type="solid")
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        thin_border_side = Side(border_style="thin", color="D3D3D3")
        grid_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = grid_border
        
        ws.row_dimensions[1].height = 26

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = grid_border
                cell.alignment = align_left

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 5, 14)

    # ==========================================
    # 1. BULLETPROOF READ EXCEL TO PYTHON
    # ==========================================
    def read_records(self, filename, headers):
        filepath = self.get_filepath(filename)
        
        if not os.path.exists(filepath):
            self.ensure_file(filename, headers)
            return []
        
        records = []
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            
            if len(rows) > 1:
                # Sanitize header names from Excel row 1
                excel_headers = [clean_val(h).lower().replace(' ', '_').replace('.', '') for h in rows[0]]
                
                for row in rows[1:]:
                    cleaned_row = [clean_val(c) for c in row]
                    if not any(cleaned_row):
                        continue
                    
                    record = {}
                    
                    # 1. Map by Excel sheet column headers
                    for i, ex_h in enumerate(excel_headers):
                        val = cleaned_row[i] if i < len(cleaned_row) else ""
                        if ex_h:
                            record[ex_h] = val

                    # 2. Positional fallback mapping for expected headers
                    for idx, exp_h in enumerate(headers):
                        norm_exp = exp_h.lower().replace(' ', '_')
                        if norm_exp not in record or not record[norm_exp]:
                            val = cleaned_row[idx] if idx < len(cleaned_row) else ""
                            record[norm_exp] = val
                    
                    # 3. Flexible Alias Fallbacks
                    if not record.get('name'):
                        record['name'] = record.get('full_name') or record.get('customer_name') or record.get('employee_name') or ""
                    if not record.get('customer'):
                        record['customer'] = record.get('customer_name') or record.get('name') or ""
                    if not record.get('tailor'):
                        record['tailor'] = record.get('assigned_tailor') or record.get('tailor_name') or record.get('staff') or ""

                    records.append(record)

        except Exception as e:
            print(f"[ExcelManager Error] Reading {filename} at {filepath}: {e}")
        
        if os.path.exists(filepath):
            self._file_mtimes[filename] = os.path.getmtime(filepath)
            
        return records

    def read_raw_matrix(self, filename):
        filepath = self.get_filepath(filename)
        if not os.path.exists(filepath):
            return [], []
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return [], []
            headers = [clean_val(h) for h in rows[0]]
            data_rows = []
            for r in rows[1:]:
                cleaned = [clean_val(c) for c in r]
                if any(cleaned):
                    data_rows.append(cleaned)
            return headers, data_rows
        except Exception as e:
            print(f"[ExcelManager Error] Reading raw matrix {filename}: {e}")
            return [], []

    # ==========================================
    # 2. WRITE & APPEND TO EXCEL WITH ERROR HANDLING
    # ==========================================
    def append_record(self, filename, headers, row_values):
        filepath = self.get_filepath(filename)
        self.ensure_file(filename, headers)
        
        try:
            wb = openpyxl.load_workbook(filepath)
            ws = wb.active
            
            clean_row = [clean_val(v) for v in row_values]
            ws.append(clean_row)
            
            self.style_worksheet(ws, headers)
            wb.save(filepath)
            self._file_mtimes[filename] = os.path.getmtime(filepath)
        except PermissionError:
            raise Exception(f"Cannot save to '{filename}'. Please close this file if it is open in Microsoft Excel!")
        except Exception as e:
            raise Exception(f"Error saving to Excel: {str(e)}")

    def write_all_records(self, filename, headers, rows_list):
        filepath = self.get_filepath(filename)
        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Data Section"
            ws.append(headers)
            
            for row in rows_list:
                clean_row = [clean_val(v) for v in row]
                ws.append(clean_row)
                
            self.style_worksheet(ws, headers)
            wb.save(filepath)
            self._file_mtimes[filename] = os.path.getmtime(filepath)
        except PermissionError:
            raise Exception(f"Cannot save to '{filename}'. Please close this file if it is open in Microsoft Excel!")
        except Exception as e:
            raise Exception(f"Error writing Excel file: {str(e)}")

    # ==========================================
    # 3. SETTINGS & CONFIG LISTS
    # ==========================================
    def load_settings_list(self, list_key, default_values):
        filepath = self.get_filepath("settings.xlsx")
        headers = ["category", "item_value"]
        
        if not os.path.exists(filepath):
            self.ensure_file("settings.xlsx", headers)
        
        items = []
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) > 1:
                for row in rows[1:]:
                    cat = clean_val(row[0])
                    val = clean_val(row[1])
                    if cat == list_key and val:
                        items.append(val)
        except Exception as e:
            print(f"[ExcelManager Error] Reading settings.xlsx: {e}")
            
        if not items:
            for dv in default_values:
                self.append_record("settings.xlsx", headers, [list_key, dv])
            items = default_values
            
        return items

    def save_settings_list(self, list_key, items_list):
        filepath = self.get_filepath("settings.xlsx")
        headers = ["category", "item_value"]
        
        other_records = []
        try:
            wb = openpyxl.load_workbook(filepath, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if len(rows) > 1:
                for row in rows[1:]:
                    cat = clean_val(row[0])
                    val = clean_val(row[1])
                    if cat and cat != list_key and val:
                        other_records.append([cat, val])
        except Exception:
            pass

        for new_item in items_list:
            if new_item.strip():
                other_records.append([list_key, new_item.strip()])

        self.write_all_records("settings.xlsx", headers, other_records)

    # ==========================================
    # 4. FILE MONITORING
    # ==========================================
    def is_file_modified(self, filename):
        filepath = self.get_filepath(filename)
        if not os.path.exists(filepath):
            return False
            
        current_mtime = os.path.getmtime(filepath)
        last_mtime = self._file_mtimes.get(filename, 0)
        
        if last_mtime != 0 and current_mtime > last_mtime:
            self._file_mtimes[filename] = current_mtime
            return True
        return False

    def ensure_file(self, filename, headers):
        filepath = self.get_filepath(filename)
        if not os.path.exists(filepath):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Data Section"
            ws.append(headers)
            self.style_worksheet(ws, headers)
            wb.save(filepath)
            self._file_mtimes[filename] = os.path.getmtime(filepath)